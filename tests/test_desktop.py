import unittest
import os
from pathlib import Path
from tempfile import TemporaryDirectory
from dataclasses import replace
from io import BytesIO
from PIL import Image
from PySide6.QtWidgets import QApplication
from PySide6.QtTest import QTest
from local_anonymizer.desktop import MainWindow
from local_anonymizer.source import LoadedSource,SourceKind,RasterPage,VisualRegion
from local_anonymizer.domain import AnonymizerError
from helpers import document

class Empty:
    fingerprint='test:1'
    def analyze(self,text,sid):return ()

class DesktopTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
    def setUp(self):
        self.w=MainWindow(detector=Empty());self.w.show();self.app.processEvents()
    def tearDown(self):
        self.idle();self.w.close();self.app.processEvents()
    def idle(self):
        for _ in range(1000):
            self.app.processEvents()
            if self.w._thread is None:return
            QTest.qWait(10)
        self.fail('job active')
    def test_batch_navigation_retains_independent_confirmations(self):
        with TemporaryDirectory() as tmp:
            a=Path(tmp)/'a.md';a.write_text('# Uno')
            b=Path(tmp)/'b.txt';b.write_text('Dos')
            self.w.open_paths((a,b));self.idle()
            self.assertIn('Uno',self.w.markdown.toPlainText())
            self.w.make_preview();self.idle();self.w.attestation.setChecked(True)
            self.w.confirm_review();self.idle();self.assertTrue(self.w.can_publish())
            first=self.w.session
            self.w.select_entry(1);self.idle()
            self.assertFalse(self.w.session.review.confirmed)
            self.assertFalse(self.w.can_publish());self.assertIsNot(first,self.w.session)
            self.w.select_entry(0);self.idle();self.assertTrue(self.w.can_publish())
    def test_region_edit_revokes_preview_and_visual_confirmation(self):
        buf=BytesIO();Image.new('RGB',(200,100),'white').save(buf,format='PNG')
        page=RasterPage(buf.getvalue(),200,100,200,100)
        source=LoadedSource(document(''),SourceKind.PDF,(page,))
        self.w.set_loaded_source(source);self.w.session.set_findings(())
        self.w.make_preview();self.idle()
        self.assertIsNotNone(self.w.entry.visual)
        self.w.add_region(VisualRegion(0,(.1,.1,.5,.5)))
        self.assertIsNone(self.w.entry.visual);self.assertIsNone(self.w.session.current_preview)
        self.assertFalse(self.w.can_publish());self.assertFalse(self.w.attestation.isChecked())
    def test_visual_confirmation_requires_viewing_pages(self):
        buf=BytesIO();Image.new('RGB',(200,100),'white').save(buf,format='PNG')
        page=RasterPage(buf.getvalue(),200,100,200,100)
        self.w.set_loaded_source(LoadedSource(document(''),SourceKind.PDF,(page,page)))
        self.w.session.set_findings(());self.w.make_preview();self.idle()
        self.w.attestation.setChecked(True);self.w.confirm_review();self.idle()
        self.assertFalse(self.w.session.review.confirmed)
        self.w.visual_output.show_page(1);self.app.processEvents()
        self.w.confirm_review();self.idle()
        self.assertTrue(self.w.can_publish())

    def test_partial_batch_exports_only_reviewed_and_preserves_existing_copy(self):
        with TemporaryDirectory() as tmp:
            root=Path(tmp);a=root/'a.txt';b=root/'b.txt';a.write_text('Uno');b.write_text('Dos')
            self.w.open_paths((a,b));self.idle();self.w.make_preview()
            self.w.attestation.setChecked(True);self.w.confirm_review();self.idle()
            dest=root/'exports';dest.mkdir()
            self.assertEqual(self.w.export_batch(dest),(1,1,0))
            original=(dest/'documento-001.txt').read_bytes()
            self.assertEqual(self.w.export_batch(dest),(0,1,1))
            self.assertEqual((dest/'documento-001.txt').read_bytes(),original)
            self.assertFalse((dest/'documento-002.txt').exists())

    def test_load_cancel_discards_late_source_and_does_not_start_analysis(self):
        import threading
        from unittest.mock import patch
        entered=threading.Event();release=threading.Event()
        def slow(*args,**kwargs):
            entered.set();release.wait(3)
            return LoadedSource(document('late'),SourceKind.TEXT)
        with TemporaryDirectory() as tmp,patch('local_anonymizer.desktop.load_source',side_effect=slow):
            p=Path(tmp)/'a.txt';p.write_text('Uno');self.w.open_paths((p,))
            for _ in range(100):
                self.app.processEvents()
                if entered.is_set():break
                QTest.qWait(5)
            self.w.cancel_work();release.set();self.idle()
            self.assertIsNone(self.w.entry.source);self.assertIsNone(self.w.session.review)
            self.assertFalse(self.w.can_publish())

    def test_visual_pdf_export_reopens_exact_images(self):
        import pymupdf
        buf=BytesIO();Image.new('RGB',(200,100),'white').save(buf,format='PNG')
        page=RasterPage(buf.getvalue(),200,100,200,100)
        self.w.set_loaded_source(LoadedSource(document(''),SourceKind.IMAGE,(page,)))
        self.w.session.set_findings(());self.w.make_preview();self.idle()
        self.w.attestation.setChecked(True);self.w.confirm_review();self.idle()
        with TemporaryDirectory() as tmp:
            path=Path(tmp)/'new.pdf';self.w.export_to(path)
            with pymupdf.open(path) as pdf:
                self.assertEqual(len(pdf),1);self.assertEqual(pdf[0].get_text(),'')
                self.assertEqual(len(pdf[0].get_images()),1)

    def test_immediate_jobs_can_chain_repeatedly_without_stuck_event_loop(self):
        with TemporaryDirectory() as tmp:
            p=Path(tmp)/'quick.txt';p.write_text('texto')
            for _ in range(15):
                self.w.open_paths((p,));self.idle();self.w.make_preview()
                self.w.attestation.setChecked(True);self.w.confirm_review();self.idle()
                self.assertTrue(self.w.can_publish())

    def test_batch_publication_cancel_preserves_completed_copy_and_review(self):
        import threading
        from unittest.mock import patch
        from local_anonymizer.export import publish_text
        entered=threading.Event();release=threading.Event()
        with TemporaryDirectory() as tmp:
            root=Path(tmp);paths=(root/'a.txt',root/'b.txt')
            for p in paths:p.write_text('texto')
            self.w.open_paths(paths);self.idle()
            for i in range(2):
                self.w.select_entry(i);self.idle();self.w.make_preview()
                self.w.attestation.setChecked(True);self.w.confirm_review();self.idle()
            def save_then_wait(*args,**kwargs):
                receipt=publish_text(*args,**kwargs);entered.set();release.wait(3);return receipt
            dest=root/'out';dest.mkdir()
            with patch('local_anonymizer.export.publish_text',side_effect=save_then_wait):
                self.w.start_batch_export(dest)
                for _ in range(100):
                    self.app.processEvents()
                    if entered.is_set():break
                    QTest.qWait(5)
                self.assertTrue(entered.is_set());self.w.cancel_work();release.set();self.idle()
            self.assertTrue((dest/'documento-001.txt').is_file())
            self.assertFalse((dest/'documento-002.txt').exists())
            self.assertIsNotNone(self.w.batch.entries[0].receipt)
            self.assertTrue(self.w.session.review.confirmed)

    @unittest.skipUnless(os.getenv('RUN_DESKTOP_INTEGRATION')=='1','explicit local Office/visual integration')
    def test_real_mixed_batch_redacts_all_formats_and_preserves_originals(self):
        import hashlib
        import re
        import pymupdf
        from docx import Document
        from docx.shared import Pt
        from PIL import ImageDraw,ImageFont
        from local_anonymizer.domain import Finding,Category,Confidentiality,Decision,Disposition,Action
        from local_anonymizer.formats import ocr_pdf
        with TemporaryDirectory() as tmp:
            root=Path(tmp);markdown=root/'source.md';markdown.write_text('# Fixture\nAna\n')
            native=root/'source.pdf'
            with pymupdf.open() as pdf:
                page=pdf.new_page(width=300,height=160);page.insert_text((40,80),'Ana',fontsize=28)
                pdf.set_metadata({'author':'Synthetic hidden author'});pdf.save(native)
            image=root/'source.png';picture=Image.new('RGB',(800,400),'white')
            ImageDraw.Draw(picture).text((100,100),'Ana',font=ImageFont.load_default(size=48),fill='black');picture.save(image)
            office=root/'source.docx';word=Document();run=word.add_paragraph().add_run('Ana');run.font.size=Pt(28)
            word.core_properties.author='Synthetic hidden author';word.save(office)
            paths=(markdown,native,image,office)
            before={p:hashlib.sha256(p.read_bytes()).hexdigest() for p in paths}
            self.w.open_paths(paths);self.idle()
            for index in range(len(paths)):
                self.w.select_entry(index);self.idle();s=self.w.session;r=s.review
                spans=tuple(re.finditer('Ana',r.document.text));self.assertTrue(spans,r.document.text)
                findings=tuple(Finding(str(i),r.document.session_id,m.start(),m.end(),Category.PERSONA,
                    Confidentiality.PII,('manual:fixture',)) for i,m in enumerate(spans))
                s.set_findings(findings);s.set_decisions(tuple(Decision(f.id,Disposition.ACCEPT,Action.REDACT) for f in findings))
                self.w.render_review();self.w.make_preview();self.idle()
                self.w.attestation.setChecked(True);self.w.confirm_review();self.idle()
                self.assertTrue(self.w.can_publish(),self.w.status.text())
            destination=root/'out';destination.mkdir();self.w.start_batch_export(destination);self.idle()
            for index,p in enumerate(paths):
                self.assertEqual(hashlib.sha256(p.read_bytes()).hexdigest(),before[p])
                ext='txt' if index==0 else 'pdf';out=destination/f'documento-{index+1:03d}.{ext}'
                self.assertTrue(out.is_file())
                if ext=='txt':self.assertNotIn('Ana',out.read_text())
                else:
                    with pymupdf.open(out) as pdf:
                        self.assertEqual(pdf[0].get_text(),'');self.assertFalse(pdf.metadata.get('author'))
                    self.assertNotIn('Ana',ocr_pdf(out.read_bytes())[0])
