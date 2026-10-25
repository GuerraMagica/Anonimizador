import io
import os
import tempfile
import time
from pathlib import Path
import warnings
import unittest
import zipfile
import struct
from unittest.mock import patch
from local_anonymizer.domain import AnonymizerError
from local_anonymizer.office import convert_docx


def package(extra=()):
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('[Content_Types].xml', '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        archive.writestr('word/document.xml', '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body/></w:document>')
        for name, data in extra:
            with warnings.catch_warnings():
                warnings.simplefilter('ignore', UserWarning)
                archive.writestr(name, data)
    return out.getvalue()


class OfficeTests(unittest.TestCase):
    def reject(self, payload, code):
        with self.assertRaises(AnonymizerError) as caught:
            convert_docx(payload)
        self.assertEqual(str(caught.exception), code)

    def test_invalid_zip(self):
        self.reject(b'private text', 'DOCX_INVALID')

    def test_content_type_declared_xml_is_inspected_despite_extension(self):
        out = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(package())) as incoming, zipfile.ZipFile(out, 'w') as archive:
            for info in incoming.infolist():
                data = incoming.read(info)
                if info.filename == '[Content_Types].xml':
                    data = data.replace(b'</Types>', b'<Override PartName="/word/header1.dat" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml"/></Types>')
                archive.writestr(info, data)
            archive.writestr('word/header1.dat', '<w:hdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:fldSimple w:instr="INCLUDETEXT &quot;/private/example&quot;"/></w:hdr>')
            archive.writestr('word/_rels/document.xml.rels', '<Relationships><Relationship Id="rHeader" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/header" Target="header1.dat"/></Relationships>')
        self.reject(out.getvalue(), 'DOCX_ACTIVE_CONTENT')

    def test_corrupt_deflate_is_sanitized(self):
        payload = bytearray(package())
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            info = archive.getinfo('word/document.xml')
            filename_len, extra_len = struct.unpack_from('<HH', payload, info.header_offset + 26)
            offset = info.header_offset + 30 + filename_len + extra_len
        payload[offset] = 0xff  # Reserved DEFLATE block type.
        self.reject(bytes(payload), 'DOCX_INVALID')

    def test_relationship_selected_xml_is_inspected_without_xml_extension(self):
        self.reject(package([
            ('word/header1.dat', '<x><fldSimple instr="INCLUDETEXT secret"/></x>'),
            ('word/_rels/document.xml.rels', '<Relationships><Relationship Type="header" Target="header1.dat"/></Relationships>'),
        ]), 'DOCX_ACTIVE_CONTENT')

    def test_termination_cleans_group_even_when_leader_already_exited(self):
        from local_anonymizer.office import _terminate
        from unittest.mock import Mock
        import signal
        process = Mock(pid=123456)
        process.poll.return_value = 0
        with patch('local_anonymizer.office.os.killpg') as killpg:
            _terminate(process)
        self.assertTrue(any(call.args == (process.pid, signal.SIGKILL) for call in killpg.call_args_list))

    @unittest.skipUnless(os.environ.get('RUN_OFFICE_INTEGRATION') == '1', 'Requires native process inspection outside nested sandbox')
    def test_termination_kills_resistant_descendant_with_live_or_exited_leader(self):
        import signal
        import subprocess
        import sys
        from local_anonymizer.office import _terminate
        for exit_leader in (False, True):
            with self.subTest(exit_leader=exit_leader), tempfile.TemporaryDirectory() as directory:
                ready = Path(directory) / 'ready'
                child = ('import os,signal,time,pathlib; '
                         'signal.signal(signal.SIGTERM, signal.SIG_IGN); '
                         f'pathlib.Path({str(ready)!r}).write_text(str(os.getpid())); '
                         'time.sleep(30)')
                leader = ('import subprocess,sys,time,pathlib; '
                          f'p=subprocess.Popen([sys.executable,"-c",{child!r}]); '
                          f'ready=pathlib.Path({str(ready)!r}); '
                          '\nwhile not ready.exists(): time.sleep(.01)\n'
                          + ('sys.exit(0)' if exit_leader else 'p.wait()'))
                process = subprocess.Popen([sys.executable, '-c', leader], start_new_session=True,
                                           stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
                try:
                    deadline = time.monotonic() + 3
                    while not ready.exists() and time.monotonic() < deadline:
                        time.sleep(.01)
                    self.assertTrue(ready.exists(), 'Child did not become ready')
                    child_pid = int(ready.read_text())
                    if exit_leader:
                        process.wait(timeout=3)
                    _terminate(process)
                    deadline = time.monotonic() + 3
                    while True:
                        state = subprocess.run(['/bin/ps', '-p', str(child_pid), '-o', 'stat='],
                                               capture_output=True, text=True).stdout.strip()
                        if not state or state.startswith('Z'):
                            break
                        if time.monotonic() > deadline:
                            self.fail('Resistant child survived group cleanup')
                        time.sleep(.01)
                finally:
                    try:
                        os.killpg(process.pid, signal.SIGKILL)
                    except ProcessLookupError:
                        pass
                    process.wait(timeout=3)

    def test_path_escape_and_duplicate(self):
        for name in ('../secret', '/absolute', 'word/../../escape', 'word\\escape', 'word/document.xml'):
            with self.subTest(name=name):
                self.reject(package([(name, b'x')]), 'DOCX_UNSAFE_PACKAGE')

    def test_expansion_and_entry_limits(self):
        with patch('local_anonymizer.office.MAX_EXPANDED_BYTES', 50):
            self.reject(package(), 'DOCX_LIMIT')
        with patch('local_anonymizer.office.MAX_ENTRIES', 1):
            self.reject(package(), 'DOCX_LIMIT')
        with patch('local_anonymizer.office.MAX_INPUT_BYTES', 1):
            self.reject(package(), 'DOCX_LIMIT')

    def test_macros_embedded_and_active_xml(self):
        for name, data in [('word/vbaProject.bin', b'x'), ('word/embeddings/oleObject1.bin', b'x'), ('word/activeX/activeX1.xml', b'<x/>'), ('word/extra.xml', b'<!DOCTYPE x [<!ENTITY a SYSTEM "file:///secret">]><x>&a;</x>')]:
            self.reject(package([(name, data)]), 'DOCX_UNSAFE_PACKAGE')

    def test_external_relationship_and_dangerous_field(self):
        self.reject(package([('word/_rels/document.xml.rels', '<Relationships><Relationship TargetMode="External" Target="https://private.invalid" Type="hyperlink"/></Relationships>')]), 'DOCX_EXTERNAL_RELATIONSHIP')
        self.reject(package([('word/extra.xml', '<x><instrText> INCLUDETEXT "/private/file" </instrText></x>')]), 'DOCX_ACTIVE_CONTENT')

    def test_cancel_before_processing(self):
        with self.assertRaises(AnonymizerError) as caught:
            convert_docx(package(), cancelled=lambda: True)
        self.assertEqual(caught.exception.code, 'CANCELLED')

    def test_engine_absent(self):
        with patch('local_anonymizer.office._find_engine', return_value=None):
            self.reject(package(), 'DOCX_ENGINE_UNAVAILABLE')

    def test_split_field_and_internal_relation_escape(self):
        self.reject(package([('word/extra.xml', '<x><instrText>INCL</instrText><instrText>UDETEXT "secret"</instrText></x>')]), 'DOCX_ACTIVE_CONTENT')
        self.reject(package([('word/_rels/document.xml.rels', '<Relationships><Relationship Target="../../outside" Type="image"/></Relationships>')]), 'DOCX_UNSAFE_PACKAGE')

    def test_containment_unavailable_never_launches_converter(self):
        with patch('local_anonymizer.office._find_engine', return_value='/engine'), patch('local_anonymizer.office.SANDBOX', '/absent-sandbox'):
            self.reject(package(), 'DOCX_CONTAINMENT_UNAVAILABLE')

    @unittest.skipUnless(os.environ.get('RUN_OFFICE_INTEGRATION') == '1', 'Requires native macOS process sandbox; run RUN_OFFICE_INTEGRATION=1 outside nested sandbox')
    def test_real_conversion_styles_table_header_metadata(self):
        import hashlib
        from docx import Document
        import pymupdf
        document = Document()
        document.add_heading('Documento sintetico', 0)
        document.add_paragraph('Acta de prueba con estilo.', 'Subtitle')
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Columna A'
        table.cell(0, 1).text = 'Columna B'
        table.cell(1, 0).text = 'Valor uno'
        table.cell(1, 1).text = 'Valor dos'
        document.sections[0].header.paragraphs[0].text = 'Cabecera sintetica'
        document.core_properties.author = 'Autor sintetico oculto'
        data = io.BytesIO()
        document.save(data)
        # Add a normal OOXML footnote without embedding objects or external links.
        enriched = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(data.getvalue())) as incoming, zipfile.ZipFile(enriched, 'w') as outgoing:
            for item in incoming.infolist():
                part = incoming.read(item)
                if item.filename == 'word/document.xml':
                    part = part.replace(b'<w:sectPr', b'<w:p><w:r><w:t>Referencia</w:t></w:r><w:r><w:footnoteReference w:id="1"/></w:r></w:p><w:sectPr')
                if item.filename == 'word/_rels/document.xml.rels':
                    part = part.replace(b'</Relationships>', b'<Relationship Id="rIdFootnote" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/></Relationships>')
                if item.filename == '[Content_Types].xml':
                    part = part.replace(b'</Types>', b'<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>')
                outgoing.writestr(item, part)
            outgoing.writestr('word/footnotes.xml', '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:id="1"><w:p><w:r><w:t>Nota sintetica visible</w:t></w:r></w:p></w:footnote></w:footnotes>')
        original = enriched.getvalue()
        digest = hashlib.sha256(original).digest()
        pdf, notes = convert_docx(original)
        self.assertEqual(hashlib.sha256(original).digest(), digest)
        self.assertTrue(any(note.startswith('DOCX_INVENTORY_METADATA_') for note in notes))
        self.assertIn('DOCX_NONVISIBLE_CONTENT_EXCLUDED', notes)
        with pymupdf.open(stream=pdf, filetype='pdf') as reopened:
            self.assertEqual(len(reopened), 1)
            text = reopened[0].get_text()
            for value in ('Cabecera sintetica', 'Documento sintetico', 'Columna A', 'Valor dos', 'Nota sintetica visible'):
                self.assertIn(value, text)
            pixmap = reopened[0].get_pixmap()
            self.assertGreater(pixmap.width, 500)
            self.assertGreater(pixmap.height, 700)

    @unittest.skipUnless(os.environ.get('RUN_OFFICE_INTEGRATION') == '1', 'Requires native process sandbox')
    def test_running_process_cancel_and_timeout(self):
        with tempfile.TemporaryDirectory() as directory:
            engine = Path(directory) / 'engine'
            engine.write_text('#!/bin/sh\nsleep 30\n')
            engine.chmod(0o700)
            started = time.monotonic()
            with patch('local_anonymizer.office._find_engine', return_value=str(engine)):
                with self.assertRaises(AnonymizerError) as caught:
                    convert_docx(package(), cancelled=lambda: time.monotonic() - started > .3)
                self.assertEqual(caught.exception.code, 'CANCELLED')
                self.assertLess(time.monotonic() - started, 5)
                with patch('local_anonymizer.office.CONVERSION_TIMEOUT', .15):
                    self.reject(package(), 'DOCX_TIMEOUT')

    @unittest.skipUnless(os.environ.get('RUN_OFFICE_INTEGRATION') == '1', 'Requires native process sandbox')
    def test_containment_denies_network_and_writes_outside_private_tree(self):
        import socket
        import sys
        with tempfile.TemporaryDirectory() as directory, socket.socket() as listener:
            listener.bind(('127.0.0.1', 0))
            listener.listen()
            port = listener.getsockname()[1]
            outside = Path(directory) / 'forbidden'
            engine = Path(directory) / 'engine'
            engine.write_text(f'''#!{sys.executable}
import socket, pathlib, sys
try:
    socket.create_connection(('127.0.0.1', {port}), timeout=1)
except OSError:
    pass
else:
    sys.exit(40)
try:
    pathlib.Path({str(outside)!r}).write_text('forbidden')
except OSError:
    pass
else:
    sys.exit(41)
pathlib.Path('input.pdf').write_bytes(b'%PDF-synthetic-containment-probe')
''')
            engine.chmod(0o700)
            with patch('local_anonymizer.office._find_engine', return_value=str(engine)):
                result, _ = convert_docx(package())
            self.assertEqual(result, b'%PDF-synthetic-containment-probe')
            self.assertFalse(outside.exists())

    def test_temporary_storage_failure_is_sanitized(self):
        with patch('local_anonymizer.office._find_engine', return_value='/engine'), patch('local_anonymizer.office.tempfile.TemporaryDirectory', side_effect=OSError('/sensitive/path')):
            self.reject(package(), 'DOCX_CONVERSION_FAILED')

    def test_inventory_never_repeats_part_names_or_text(self):
        from local_anonymizer.office import _validate
        notes = _validate(package([('docProps/private-author.xml', '<metadata>Secret Person</metadata>'), ('_xmlsignatures/private-signature.xml', '<signature/>'), ('word/comments.xml', '<comments/>'), ('customXml/private.xml', '<private/>')]), lambda: False)
        joined = '\n'.join(notes)
        self.assertNotIn('Secret Person', joined)
        self.assertNotIn('private', joined)
        self.assertIn('DOCX_INVENTORY_SIGNATURES_1', joined)
        self.assertIn('DOCX_INVENTORY_COMMENTS_1', joined)
